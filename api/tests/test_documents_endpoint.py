"""POST /documents/upload -- camada de API (roteamento, validação, sessão),
mesmo escopo/estilo de test_endpoints.py (não re-testa o motor de extração
em si, já coberto por docksmith/tests/test_document_loader.py).
"""

from __future__ import annotations

import io

import fitz
from docx import Document as DocxDocument

from api import auth, config, r2_storage
from api.main import app
from api.tests.test_r2_storage import _FakeR2Client


def _login_as(user_id: str):
    app.dependency_overrides[auth.get_current_user] = lambda: {
        "token": "t",
        "user": {"id": user_id, "email": f"{user_id}@docksmith.local"},
    }


def _pdf_bytes(pages_text: list[str]) -> bytes:
    pdf = fitz.open()
    for text in pages_text:
        page = pdf.new_page()
        page.insert_text((72, 72), text)
    data = pdf.tobytes()
    pdf.close()
    return data


def _docx_bytes(paragraphs: list[str]) -> bytes:
    doc = DocxDocument()
    for text in paragraphs:
        doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _multi_chapter_pdf_bytes(chapter_page_counts: list[int]) -> bytes:
    """Mesmo formato de docksmith/tests/test_document_loader.py::
    _build_multi_chapter_pdf_bytes -- reescrito aqui (não importado de lá)
    pra não criar dependência entre as 2 suítes: aquela testa o motor de
    extração/divisão em si, esta testa só a camada de API por cima dele."""
    pdf = fitz.open()
    cover = pdf.new_page()
    cover.insert_text((72, 200), "Título Gigante da Capa", fontsize=48)
    for chapter_index, page_count in enumerate(chapter_page_counts, start=1):
        for page_in_chapter in range(page_count):
            page = pdf.new_page()
            if page_in_chapter == 0:
                page.insert_text((72, 72), f"CHAPTER {chapter_index}", fontsize=18)
                y = 110
            else:
                y = 72
            for line in range(8):
                page.insert_text(
                    (72, y),
                    f"Texto de corpo real, capítulo {chapter_index}, página {page_in_chapter + 1}, linha {line}.",
                    fontsize=11,
                )
                y += 16
    data = pdf.tobytes()
    pdf.close()
    return data


def _docx_chapters_bytes(chapter_subsection_counts: list[int]) -> bytes:
    doc = DocxDocument()
    for chapter_index, subsection_count in enumerate(chapter_subsection_counts, start=1):
        doc.add_heading(f"Capítulo {chapter_index}", level=1)
        for subsection_index in range(1, subsection_count + 1):
            doc.add_heading(f"Capítulo {chapter_index}.{subsection_index}", level=2)
            doc.add_paragraph(f"Parágrafo da subseção {subsection_index} do capítulo {chapter_index}.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_upload_pdf_creates_a_collection(client):
    _login_as("user-1")
    resp = client.post(
        "/documents/upload",
        files={"file": ("manual.pdf", _pdf_bytes(["conteúdo página um", "conteúdo página dois"]), "application/pdf")},
        data={"collection_name": "manual-pdf"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["collection_name"] == "manual-pdf"
    assert body["document_count"] == 2
    assert body["session_id"]


def test_upload_docx_creates_a_collection(client):
    _login_as("user-1")
    resp = client.post(
        "/documents/upload",
        files={
            "file": (
                "guia.docx",
                _docx_bytes(["Primeiro parágrafo real do documento."]),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"collection_name": "guia-docx"},
    )
    assert resp.status_code == 200
    assert resp.json()["document_count"] == 1


def test_upload_requires_auth(client):
    resp = client.post(
        "/documents/upload",
        files={"file": ("manual.pdf", _pdf_bytes(["x"]), "application/pdf")},
        data={"collection_name": "c"},
    )
    assert resp.status_code == 401


def test_upload_rejects_unsupported_extension(client):
    _login_as("user-1")
    resp = client.post(
        "/documents/upload",
        files={"file": ("planilha.xlsx", b"conteudo qualquer", "application/vnd.ms-excel")},
        data={"collection_name": "c"},
    )
    assert resp.status_code == 422
    assert "não suportado" in resp.json()["detail"]


def test_upload_rejects_scanned_pdf_with_a_clear_error(client):
    _login_as("user-1")
    pdf = fitz.open()
    pdf.new_page()  # página em branco, sem texto -- simula um scan
    data = pdf.tobytes()
    pdf.close()

    resp = client.post(
        "/documents/upload",
        files={"file": ("scan.pdf", data, "application/pdf")},
        data={"collection_name": "c"},
    )
    assert resp.status_code == 422
    assert "escaneado" in resp.json()["detail"]


def test_upload_rejects_empty_file(client):
    _login_as("user-1")
    resp = client.post(
        "/documents/upload",
        files={"file": ("vazio.pdf", b"", "application/pdf")},
        data={"collection_name": "c"},
    )
    assert resp.status_code == 422


def test_upload_rejects_file_over_size_limit(client, monkeypatch):
    _login_as("user-1")
    monkeypatch.setattr(config, "DOCSMITH_MAX_FILE_SIZE_MB", 0)  # qualquer arquivo não-vazio excede 0MB
    # api/routers/documents.py lê `_MAX_FILE_BYTES` uma vez na importação do
    # módulo -- reaplica o mesmo patch no atributo já calculado do router.
    import api.routers.documents as documents_router
    monkeypatch.setattr(documents_router, "_MAX_FILE_BYTES", 0)

    resp = client.post(
        "/documents/upload",
        files={"file": ("manual.pdf", _pdf_bytes(["x"]), "application/pdf")},
        data={"collection_name": "c"},
    )
    assert resp.status_code == 413
    # 2026-09-10: achado real testando ao vivo pelo Hub -- o texto do erro
    # vazava o nome literal da variável de ambiente interna pro usuário
    # final ("...limite de 0MB (DOCSMITH_MAX_FILE_SIZE_MB)."). Confirma que
    # não vaza mais, sem exigir uma string exata (frágil a reformulação).
    assert "DOCSMITH_MAX_FILE_SIZE_MB" not in resp.json()["detail"]


# ===== Limite de páginas: estruturado + truncamento opcional (2026-09-10,
# pedido explícito do usuário testando um livro real de 446 páginas pelo
# Hub) =====


def test_upload_rejects_file_over_page_limit_with_structured_error(client, monkeypatch):
    monkeypatch.setattr(config, "DOCSMITH_MAX_PAGES", 3)
    _login_as("user-1")

    resp = client.post(
        "/documents/upload",
        files={"file": ("livro.pdf", _pdf_bytes([f"página {i}" for i in range(5)]), "application/pdf")},
        data={"collection_name": "livro"},
    )
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    # Estruturado (dict), não só uma string -- é isso que permite o
    # frontend oferecer "processar só as primeiras N páginas" em vez de só
    # mostrar um erro sem saída.
    assert detail["error"] == "page_limit_exceeded"
    assert detail["actual_count"] == 5
    assert detail["max_pages"] == 3
    assert detail["unit"] == "páginas"
    # Mesmo achado de vazamento de nome de variável interna do teste acima,
    # confirmado aqui também.
    assert "DOCSMITH_MAX_PAGES" not in detail["message"]


def test_upload_with_truncate_accepts_a_file_over_the_page_limit(client, monkeypatch):
    monkeypatch.setattr(config, "DOCSMITH_MAX_PAGES", 3)
    _login_as("user-1")

    resp = client.post(
        "/documents/upload",
        files={"file": ("livro.pdf", _pdf_bytes([f"página {i}" for i in range(5)]), "application/pdf")},
        data={"collection_name": "livro", "truncate": "true"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["document_count"] == 3  # só as 3 primeiras páginas, não as 5
    assert body["truncated"] is True


def test_upload_within_page_limit_reports_truncated_false_by_default(client):
    """`truncated` só é True quando o campo `truncate` foi mandado -- o
    caminho comum (documento dentro do limite) não muda de comportamento."""
    _login_as("user-1")
    resp = client.post(
        "/documents/upload",
        files={"file": ("manual.pdf", _pdf_bytes(["x"]), "application/pdf")},
        data={"collection_name": "c"},
    )
    assert resp.status_code == 200
    assert resp.json()["truncated"] is False


def test_upload_streamed_in_multiple_chunks_reconstructs_the_file_correctly(client, monkeypatch):
    """2026-09-02, 4ª rodada: o upload passou a ler em pedaços de 1MB (pra
    nunca bufferizar um arquivo inteiro de até 1GB de uma vez só antes de
    checar o tamanho). Um PDF de teste normal cabe todo num único pedaço
    de 1MB -- pra provar de verdade que a MONTAGEM de vários pedaços
    reconstrói o arquivo original byte a byte (não só "não deu erro"),
    este teste reduz o tamanho do pedaço bem abaixo do tamanho real do
    arquivo, forçando várias iterações do loop."""
    import api.routers.documents as documents_router

    monkeypatch.setattr(documents_router, "_UPLOAD_CHUNK_BYTES", 37)  # bem menor que o PDF de teste
    _login_as("user-1")

    pdf_bytes = _pdf_bytes(["conteúdo real espalhado por várias páginas para garantir bytes suficientes"] * 3)
    assert len(pdf_bytes) > 37 * 5  # confirma que o teste de fato força múltiplas iterações

    resp = client.post(
        "/documents/upload",
        files={"file": ("grande.pdf", pdf_bytes, "application/pdf")},
        data={"collection_name": "streamed"},
    )
    assert resp.status_code == 200
    assert resp.json()["document_count"] == 3  # as 3 páginas foram extraídas normalmente


def test_uploading_a_second_file_into_the_same_collection_appends_documents(client):
    _login_as("user-1")
    first = client.post(
        "/documents/upload",
        files={"file": ("parte1.pdf", _pdf_bytes(["conteúdo da primeira parte"]), "application/pdf")},
        data={"collection_name": "colecao-combinada"},
    )
    session_id = first.json()["session_id"]

    second = client.post(
        "/documents/upload",
        files={"file": ("parte2.pdf", _pdf_bytes(["conteúdo da segunda parte"]), "application/pdf")},
        data={"collection_name": "colecao-combinada", "session_id": session_id},
    )

    assert second.status_code == 200
    assert second.json()["document_count"] == 2  # 1 da primeira parte + 1 da segunda, não substituiu


def test_uploaded_document_labels_are_stored_alongside_the_collection(client):
    """Confirma o achado real desta rodada: `session["collection_labels"]`
    passa a existir e carrega o rótulo real por documento (ex.: "arquivo.pdf
    - página 1"), consumido depois por RAGService.load_collection via
    api/routers/chat.py."""
    _login_as("user-1")
    from api import sessions

    resp = client.post(
        "/documents/upload",
        files={"file": ("relatorio.pdf", _pdf_bytes(["primeira página", "segunda página"]), "application/pdf")},
        data={"collection_name": "relatorio"},
    )
    session_id = resp.json()["session_id"]
    session = sessions.get_session(session_id, "user-1")

    labels = session["collection_labels"]["relatorio"]
    assert labels == ["relatorio.pdf - página 1", "relatorio.pdf - página 2"]


def test_uploaded_document_structure_is_stored_alongside_the_collection(client):
    """3ª rodada (prompt-mestre §13/§16): `session["collection_structure"]`
    passa a existir, paralela a `collection_labels`, com page_start/
    page_end reais por página de um PDF -- é isso que
    api/routers/chat.py repassa pra RAGService.load_collection depois."""
    from api import sessions

    _login_as("user-1")

    resp = client.post(
        "/documents/upload",
        files={"file": ("relatorio2.pdf", _pdf_bytes(["primeira página", "segunda página"]), "application/pdf")},
        data={"collection_name": "relatorio2"},
    )
    session_id = resp.json()["session_id"]
    session = sessions.get_session(session_id, "user-1")

    structure = session["collection_structure"]["relatorio2"]
    assert len(structure) == 2
    assert structure[0]["page_start"] == 1 and structure[0]["page_end"] == 1
    assert structure[1]["page_start"] == 2 and structure[1]["page_end"] == 2


# ===== Divisão inteligente em partes (Fase A, 2026-09-10, "PROMPT DE
# EVOLUÇÃO"/"OBSERVAÇÕES FINAIS" -- substitui truncamento como único
# caminho pra documento acima de DOCSMITH_MAX_PAGES) =====


def test_upload_split_requires_auth(client):
    resp = client.post(
        "/documents/upload/split",
        files={"file": ("livro.pdf", _pdf_bytes(["x"]), "application/pdf")},
        data={"collection_name": "livro"},
    )
    assert resp.status_code == 401


def test_upload_split_rejects_unsupported_extension(client):
    _login_as("user-1")
    resp = client.post(
        "/documents/upload/split",
        files={"file": ("planilha.xlsx", b"conteudo qualquer", "application/vnd.ms-excel")},
        data={"collection_name": "c"},
    )
    assert resp.status_code == 422
    assert "não suportado" in resp.json()["detail"]


def test_upload_split_creates_one_collection_per_part_covering_the_whole_document(client, monkeypatch):
    monkeypatch.setattr(config, "DOCSMITH_MAX_PAGES", 10)
    _login_as("user-1")

    resp = client.post(
        "/documents/upload/split",
        files={"file": ("livro.pdf", _multi_chapter_pdf_bytes([5, 5, 30, 5]), "application/pdf")},
        data={"collection_name": "livro-grande"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["base_collection_name"] == "livro-grande"
    assert len(body["parts"]) > 1  # 30 páginas do capítulo 3 forçam mais de 1 parte

    total_pages = sum(part["page_count"] for part in body["parts"])
    assert total_pages == 1 + 5 + 5 + 30 + 5  # capa + 4 capítulos, nada perdido

    for part in body["parts"]:
        assert part["page_count"] <= 10  # nenhuma parte excede o limite configurado
        assert part["unit"] == "páginas"


def test_upload_split_each_part_is_queryable_as_its_own_full_collection(client, monkeypatch):
    """As coleções criadas pelo split precisam funcionar exatamente como
    qualquer coleção normal (labels/structure preenchidos) -- o resto do
    sistema não pode tratar isso como um caso especial escondido."""
    from api import sessions

    monkeypatch.setattr(config, "DOCSMITH_MAX_PAGES", 3)
    _login_as("user-1")

    resp = client.post(
        "/documents/upload/split",
        files={"file": ("livro.pdf", _multi_chapter_pdf_bytes([2, 2]), "application/pdf")},
        data={"collection_name": "livro-pequeno"},
    )
    assert resp.status_code == 200
    body = resp.json()
    session = sessions.get_session(body["session_id"], "user-1")

    for part in body["parts"]:
        name = part["collection_name"]
        assert name in session["collections"]
        assert name in session["collection_labels"]
        assert name in session["collection_structure"]
        assert len(session["collections"][name]) == part["document_count"]
        assert len(session["collection_labels"][name]) == part["document_count"]


def test_upload_split_reports_chunk_and_size_estimates_per_part(client, monkeypatch):
    """Fase B do plano ("metadados completos por parte") -- confirma que
    a resposta já reaproveita `resource_estimate` de verdade, não zeros
    fixos."""
    monkeypatch.setattr(config, "DOCSMITH_MAX_PAGES", 10)
    _login_as("user-1")

    resp = client.post(
        "/documents/upload/split",
        files={"file": ("livro.pdf", _multi_chapter_pdf_bytes([5, 5, 30, 5]), "application/pdf")},
        data={"collection_name": "livro-estimado"},
    )
    assert resp.status_code == 200
    for part in resp.json()["parts"]:
        assert part["estimated_chunks"] > 0
        # `estimated_size_mb` é arredondado a 2 casas (api/resource_
        # estimate.py) -- textos sintéticos de teste são pequenos o
        # bastante pra arredondar pra 0.0 de forma legítima; só
        # `estimated_chunks` (nunca arredondado a 0 pra conteúdo
        # não-vazio) é o sinal real de "a estimativa rodou de verdade".
        assert part["estimated_size_mb"] >= 0


def test_upload_split_docx_uses_section_as_the_position_unit(client, monkeypatch):
    """DOCX não tem "página" real -- a unidade reportada precisa ser
    "seções", diferente do PDF ("páginas")."""
    monkeypatch.setattr(config, "DOCSMITH_MAX_PAGES", 3)
    _login_as("user-1")

    resp = client.post(
        "/documents/upload/split",
        files={
            "file": (
                "guia.docx",
                _docx_chapters_bytes([2, 2]),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"collection_name": "guia-dividido"},
    )
    assert resp.status_code == 200
    parts = resp.json()["parts"]
    assert len(parts) >= 1
    for part in parts:
        assert part["unit"] == "seções"


def test_upload_split_document_within_the_limit_still_produces_a_single_part(client, monkeypatch):
    monkeypatch.setattr(config, "DOCSMITH_MAX_PAGES", 300)
    _login_as("user-1")

    resp = client.post(
        "/documents/upload/split",
        files={"file": ("pequeno.pdf", _multi_chapter_pdf_bytes([2]), "application/pdf")},
        data={"collection_name": "pequeno"},
    )
    assert resp.status_code == 200
    assert len(resp.json()["parts"]) == 1


def test_upload_split_saves_a_single_backup_of_the_original_file(client, monkeypatch):
    """Fase D do plano ("preservação do arquivo original, 1 backup só") --
    confirma que dividir em N partes não cria N backups no R2, só o
    arquivo original de sempre."""
    fake = _configure_r2(monkeypatch)
    monkeypatch.setattr(config, "DOCSMITH_MAX_PAGES", 10)
    _login_as("user-1")

    resp = client.post(
        "/documents/upload/split",
        files={"file": ("livro.pdf", _multi_chapter_pdf_bytes([5, 5, 30, 5]), "application/pdf")},
        data={"collection_name": "livro-backup"},
    )
    assert resp.status_code == 200
    assert len(resp.json()["parts"]) > 1  # de fato foi dividido em várias partes

    session_id = resp.json()["session_id"]
    objects = r2_storage.list_session_objects(user_id="user-1", session_id=session_id)
    assert len(objects) == 1  # 1 backup só, não 1 por parte
    assert objects[0]["filename"] == "livro.pdf"
    assert fake.objects


# ===== R2 (2026-09-02, 4ª rodada): cópia de segurança no upload + restauração =====


def _configure_r2(monkeypatch) -> _FakeR2Client:
    monkeypatch.setattr(config, "R2_ACCOUNT_ID", "acc123")
    monkeypatch.setattr(config, "R2_ACCESS_KEY_ID", "key")
    monkeypatch.setattr(config, "R2_SECRET_ACCESS_KEY", "secret")
    monkeypatch.setattr(config, "R2_BUCKET_NAME", "docksmith-bucket")
    fake = _FakeR2Client()
    monkeypatch.setattr(r2_storage, "_get_client", lambda: fake)
    return fake


def test_upload_without_r2_configured_still_succeeds_normally(client, monkeypatch):
    """R2 nunca é pré-requisito -- sem as 4 variáveis, o upload funciona
    exatamente como sempre funcionou (a chamada de fundo pro R2 vira
    no-op, `is_configured()` barra antes de qualquer tentativa de rede)."""
    monkeypatch.setattr(config, "R2_ACCOUNT_ID", None)
    _login_as("user-1")
    resp = client.post(
        "/documents/upload",
        files={"file": ("manual.pdf", _pdf_bytes(["conteúdo"]), "application/pdf")},
        data={"collection_name": "sem-r2"},
    )
    assert resp.status_code == 200


def test_upload_with_r2_configured_saves_a_real_backup_copy(client, monkeypatch):
    fake = _configure_r2(monkeypatch)
    _login_as("user-1")
    resp = client.post(
        "/documents/upload",
        files={"file": ("manual.pdf", _pdf_bytes(["conteúdo real do manual"]), "application/pdf")},
        data={"collection_name": "com-r2"},
    )
    assert resp.status_code == 200
    session_id = resp.json()["session_id"]

    # BackgroundTasks roda depois da resposta, mas dentro do mesmo
    # request síncrono do TestClient -- já deve ter rodado quando
    # chegamos aqui.
    objects = r2_storage.list_session_objects(user_id="user-1", session_id=session_id)
    assert len(objects) == 1
    assert objects[0]["collection_name"] == "com-r2"
    assert objects[0]["filename"] == "manual.pdf"
    assert fake.objects  # o "arquivo" de fato está no armazenamento falso


def test_restore_reports_session_still_alive_when_nothing_expired(client, monkeypatch):
    _configure_r2(monkeypatch)
    _login_as("user-1")
    upload = client.post(
        "/documents/upload",
        files={"file": ("manual.pdf", _pdf_bytes(["x"]), "application/pdf")},
        data={"collection_name": "c"},
    )
    session_id = upload.json()["session_id"]

    resp = client.post(f"/sessions/{session_id}/restore")
    assert resp.status_code == 200
    body = resp.json()
    assert body["restored"] is False
    assert body["reason"] == "session_still_alive"
    assert body["collections"] == ["c"]


def test_restore_reports_nothing_to_restore_for_an_unknown_session(client, monkeypatch):
    _configure_r2(monkeypatch)
    _login_as("user-1")
    resp = client.post("/sessions/inexistente-123/restore")
    assert resp.status_code == 200
    body = resp.json()
    assert body["restored"] is False
    assert body["reason"] == "nothing_to_restore"
    assert body["collections"] == []


def test_restore_without_r2_configured_reports_nothing_to_restore(client, monkeypatch):
    monkeypatch.setattr(config, "R2_ACCOUNT_ID", None)
    _login_as("user-1")
    resp = client.post("/sessions/qualquer/restore")
    assert resp.status_code == 200
    assert resp.json()["reason"] == "nothing_to_restore"


def test_restore_brings_the_documents_back_after_the_in_memory_session_is_gone(client, monkeypatch):
    """O teste central desta feature: simula exatamente o cenário real --
    upload normal, servidor reinicia (sessão em memória some), cliente
    chama /restore com o MESMO session_id que já tinha, e os documentos
    voltam a existir numa sessão nova com o mesmo id, prontos pra
    /chat funcionar de novo sem o usuário precisar reenviar o arquivo."""
    from api import sessions

    _configure_r2(monkeypatch)
    _login_as("user-1")
    upload = client.post(
        "/documents/upload",
        files={"file": ("relatorio.pdf", _pdf_bytes(["primeira página real", "segunda página real"]), "application/pdf")},
        data={"collection_name": "relatorio"},
    )
    session_id = upload.json()["session_id"]

    # Simula o reinício do servidor: a sessão em memória desaparece, mas
    # o backup no R2 (armazenamento à parte) continua existindo.
    sessions._sessions.clear()
    assert sessions.get_session(session_id, "user-1") is None

    resp = client.post(f"/sessions/{session_id}/restore")
    assert resp.status_code == 200
    body = resp.json()
    assert body["restored"] is True
    assert body["collections"] == ["relatorio"]

    # A sessão restaurada usa o MESMO session_id -- o cliente não precisa
    # trocar de id, só continuar usando o que já tinha.
    restored_session = sessions.get_session(session_id, "user-1")
    assert restored_session is not None
    assert restored_session["collections"]["relatorio"] == ["primeira página real", "segunda página real"]
    # Estrutura (capítulo/página) também sobrevive à restauração, não só o texto.
    assert restored_session["collection_structure"]["relatorio"][0]["page_start"] == 1


def test_restore_never_leaks_another_users_documents(client, monkeypatch):
    """Mesmo achando o session_id certo, restaurar como uma pessoa
    DIFERENTE da que fez o upload original não traz nada de volta -- ver
    test_r2_storage.py::test_a_different_user_id_never_sees_another_users_session_objects
    pra a mesma garantia no nível mais baixo."""
    from api import sessions

    _configure_r2(monkeypatch)
    _login_as("dono-real")
    upload = client.post(
        "/documents/upload",
        files={"file": ("confidencial.pdf", _pdf_bytes(["informação sensível"]), "application/pdf")},
        data={"collection_name": "privado"},
    )
    session_id = upload.json()["session_id"]
    sessions._sessions.clear()  # simula reinício

    _login_as("outra-pessoa")
    resp = client.post(f"/sessions/{session_id}/restore")
    assert resp.status_code == 200
    body = resp.json()
    assert body["restored"] is False
    assert body["reason"] == "nothing_to_restore"
