"""Testes de docksmith/service/document_loader.py (prompt-mestre "Docksmith"
§6/§8/§10, 2ª rodada da Fase I) -- gera PDF/DOCX reais em memória com as
próprias bibliotecas de extração (fitz/python-docx), evitando fixtures
binárias versionadas no repositório.
"""

from __future__ import annotations

import io

import fitz
import pytest
from docx import Document as DocxDocument

from docksmith.service.document_loader import (
    PageLimitExceededError,
    UnsupportedDocumentError,
    load_docx,
    load_document,
    load_pdf,
    split_docx_into_parts,
    split_document_into_parts,
    split_pdf_into_parts,
)


def _build_pdf_bytes(pages_text: list[str]) -> bytes:
    pdf = fitz.open()
    for text in pages_text:
        page = pdf.new_page()
        page.insert_text((72, 72), text)
    data = pdf.tobytes()
    pdf.close()
    return data


def _build_book_with_cover_and_chapters_pdf_bytes() -> bytes:
    """Reproduz o formato real que expôs o bug de detecção de capítulo
    (2026-09-10, achado testando um livro real de 446 páginas pelo Hub,
    não hipotético): capa com título em fonte GIGANTE (maior que qualquer
    heading real do livro, só aparece 1 vez) + vários capítulos reais,
    cada um com um rótulo curto tipo "CHAPTER N" (fonte média, se repete
    1x por capítulo) e um subtítulo de conteúdo dentro dele (fonte menor
    que o rótulo do capítulo, mas se repete MUITAS vezes ao longo do
    livro -- múltiplos subtítulos por capítulo). É a combinação exata que
    fazia as 2 primeiras tentativas de correção escolherem o tamanho de
    fonte errado."""
    pdf = fitz.open()

    cover = pdf.new_page()
    cover.insert_text((72, 200), "Título Gigante da Capa", fontsize=48)

    for chapter_num in range(1, 4):
        page = pdf.new_page()
        page.insert_text((72, 72), f"CHAPTER {chapter_num}", fontsize=18)
        # Vários subtítulos de conteúdo por capítulo (fonte menor que o
        # rótulo do capítulo, mas repetidos muitas vezes no total) -- é
        # isso que faz esse tamanho "vencer" por frequência bruta se o
        # critério não for específico o bastante.
        y = 110
        for section_num in range(1, 4):
            page.insert_text((72, y), f"Subtópico {section_num}", fontsize=14)
            y += 25
            # Corpo de texto real embaixo de cada subtítulo -- sem isso, a
            # MEDIANA de tamanho de fonte do documento (que decide o que é
            # "heading" vs "corpo") fica distorcida pela pouca quantidade
            # de texto na fixture, e nem os subtítulos são detectados como
            # heading. Um livro real tem muito mais corpo que heading;
            # esta fixture precisa do mesmo formato pra ser um teste real.
            for line in range(6):
                page.insert_text((72, y), f"Texto de corpo real da seção, linha {line}.", fontsize=11)
                y += 16

    data = pdf.tobytes()
    pdf.close()
    return data


def _build_multi_chapter_pdf_bytes(chapter_page_counts: list[int]) -> bytes:
    """PDF sintético com 1 página de capa + N capítulos, cada um com o nº
    de páginas pedido em `chapter_page_counts` (índice 0 = Capítulo 1) --
    corpo de texto real em toda página, mesma técnica já validada em
    `_build_book_with_cover_and_chapters_pdf_bytes` (evita a mediana de
    fonte distorcida por página sem corpo). Só a 1ª página de cada
    capítulo carrega o rótulo "CHAPTER N" (como num livro real), as
    demais só têm corpo -- a herança por página de `_assign_chapter_and_
    section` cobre o resto."""
    pdf = fitz.open()
    cover = pdf.new_page()
    cover.insert_text((72, 200), "Título Gigante da Capa", fontsize=48)

    for chapter_index, page_count in enumerate(chapter_page_counts, start=1):
        for page_in_chapter in range(page_count):
            page = pdf.new_page()
            if page_in_chapter == 0:
                page.insert_text((72, 72), f"CHAPTER {chapter_index}", fontsize=18)
                y = 110
            else:
                y = 72
            for line in range(8):
                page.insert_text(
                    (72, y),
                    f"Texto de corpo real, capítulo {chapter_index}, página {page_in_chapter + 1}, linha {line}.",
                    fontsize=11,
                )
                y += 16

    data = pdf.tobytes()
    pdf.close()
    return data


def _build_pdf_without_chapter_labels_bytes(page_count: int) -> bytes:
    """PDF sintético com um "título" por página (fonte maior que o corpo,
    detectável como heading) mas que NÃO segue o padrão de rótulo de
    capítulo real (ex. "Seção 1: assunto qualquer", não "CHAPTER 1") --
    caso de "sem estrutura de capítulo confiável", usado pra confirmar o
    fallback honesto de janela de página pura."""
    pdf = fitz.open()
    for page_num in range(page_count):
        page = pdf.new_page()
        page.insert_text((72, 72), f"Seção {page_num + 1}: assunto qualquer", fontsize=16)
        y = 110
        for line in range(8):
            page.insert_text((72, y), f"Texto de corpo real da página {page_num + 1}, linha {line}.", fontsize=11)
            y += 16
    data = pdf.tobytes()
    pdf.close()
    return data


def _build_docx_with_chapters_bytes(chapter_subsection_counts: list[int]) -> bytes:
    """DOCX sintético com N capítulos (Heading 1), cada um com o nº de
    subseções (Heading 2) pedido em `chapter_subsection_counts` -- cada
    subseção carrega 1 parágrafo, formando sua própria `LoadedDocument`
    (a unidade real de divisão pro DOCX, equivalente a "1 página" no PDF
    -- um Heading 1 sozinho sem parágrafo próprio não vira `LoadedDocument`,
    só marca `chapter` pras subseções seguintes). Usado pra testar
    `split_docx_into_parts` com capítulos de tamanhos desiguais."""
    doc = DocxDocument()
    for chapter_index, subsection_count in enumerate(chapter_subsection_counts, start=1):
        doc.add_heading(f"Capítulo {chapter_index}", level=1)
        for subsection_index in range(1, subsection_count + 1):
            doc.add_heading(f"Capítulo {chapter_index}.{subsection_index}", level=2)
            doc.add_paragraph(f"Parágrafo da subseção {subsection_index} do capítulo {chapter_index}.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_scanned_pdf_bytes(page_count: int = 1) -> bytes:
    """PDF com páginas em branco (sem nenhum texto inserido) -- simula um
    scan (só imagem, sem camada de texto), o caso que o loader precisa
    detectar e recusar com uma mensagem clara em vez de OCR silencioso."""
    pdf = fitz.open()
    for _ in range(page_count):
        pdf.new_page()
    data = pdf.tobytes()
    pdf.close()
    return data


def _build_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Introdução", level=1)
    doc.add_paragraph("Este é o parágrafo introdutório do documento de teste.")
    doc.add_heading("Configuração", level=1)
    doc.add_paragraph("Primeiro passo da configuração.")
    doc.add_paragraph("Segundo passo da configuração.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Campo"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "Timeout"
    table.cell(1, 1).text = "30s"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class TestLoadPdf:
    def test_extracts_one_document_per_page_in_order(self):
        pdf_bytes = _build_pdf_bytes(["Conteúdo da página um.", "Conteúdo da página dois."])

        documents = load_pdf(pdf_bytes, "manual.pdf", max_pages=300)

        assert len(documents) == 2
        assert "página um" in documents[0].text
        assert "página dois" in documents[1].text
        assert documents[0].source_label == "manual.pdf - página 1"
        assert documents[1].source_label == "manual.pdf - página 2"

    def test_scanned_pdf_without_text_raises_a_clear_error(self):
        pdf_bytes = _build_scanned_pdf_bytes(page_count=2)

        with pytest.raises(UnsupportedDocumentError, match="escaneado"):
            load_pdf(pdf_bytes, "scan.pdf", max_pages=300)

    def test_pdf_over_max_pages_is_rejected(self):
        pdf_bytes = _build_pdf_bytes([f"página {i}" for i in range(5)])

        with pytest.raises(UnsupportedDocumentError, match="limite"):
            load_pdf(pdf_bytes, "grande.pdf", max_pages=3)

    def test_corrupted_pdf_raises_a_clear_error(self):
        with pytest.raises(UnsupportedDocumentError):
            load_pdf(b"isso nao e um pdf de verdade", "quebrado.pdf", max_pages=300)

    def test_chapter_detection_is_not_confused_by_a_large_cover_title(self):
        """Bug real corrigido (2026-09-10, achado testando um livro real de
        446 páginas pelo Hub) -- a capa (fonte gigante, aparece 1x só) não
        pode "sequestrar" a detecção de capítulo pro resto do livro
        inteiro. Cada página de conteúdo precisa carregar o `chapter`
        certo (ex. "CHAPTER 2" pras páginas do capítulo 2, não o texto da
        capa nem `None`)."""
        pdf_bytes = _build_book_with_cover_and_chapters_pdf_bytes()

        documents = load_pdf(pdf_bytes, "livro.pdf", max_pages=300)

        # página 1 = capa (sem chapter, ainda não viu nenhum heading de
        # capítulo); página 2 = CHAPTER 1; página 3 = CHAPTER 2; página 4 = CHAPTER 3.
        by_page = {doc.page_start: doc.chapter for doc in documents}
        assert by_page[1] is None  # capa, antes de qualquer capítulo
        assert by_page[2] == "CHAPTER 1"
        assert by_page[3] == "CHAPTER 2"
        assert by_page[4] == "CHAPTER 3"

    def test_chapter_detection_is_not_confused_by_more_frequent_content_subtitles(self):
        """Mesmo achado acima, ângulo diferente: os subtítulos de conteúdo
        (fonte menor, mas 3x mais frequentes que o rótulo de capítulo
        nesta fixture) não podem "vencer" a escolha de qual tamanho de
        fonte representa capítulo -- confirma que o critério é o padrão
        do TEXTO (rótulo de capítulo de verdade), não frequência bruta."""
        pdf_bytes = _build_book_with_cover_and_chapters_pdf_bytes()

        documents = load_pdf(pdf_bytes, "livro.pdf", max_pages=300)

        chapters_seen = {doc.chapter for doc in documents if doc.chapter}
        assert chapters_seen == {"CHAPTER 1", "CHAPTER 2", "CHAPTER 3"}
        # Os subtítulos viraram `section`, o nível certo -- não vazaram
        # pro campo `chapter`. Cada `LoadedDocument` é 1 POR PÁGINA (não
        # por subtítulo), então `section` reflete o último subtítulo
        # visto até o fim daquela página -- aqui, sempre "Subtópico 3"
        # (o 3º e último inserido em cada página de capítulo).
        sections_seen = {doc.section for doc in documents if doc.section}
        assert sections_seen == {"Subtópico 3"}

    def test_pdf_over_max_pages_raises_page_limit_exceeded_with_real_numbers(self):
        """2026-09-10: a mensagem genérica virou uma exceção estruturada --
        confirma que os números batem de verdade (não só que "algum erro"
        foi levantado), já que é isso que a API usa pra montar a resposta
        que o frontend consome pra oferecer o truncamento."""
        pdf_bytes = _build_pdf_bytes([f"página {i}" for i in range(5)])

        with pytest.raises(PageLimitExceededError) as exc_info:
            load_pdf(pdf_bytes, "grande.pdf", max_pages=3)

        assert exc_info.value.actual_count == 5
        assert exc_info.value.max_pages == 3
        assert exc_info.value.unit == "páginas"

    def test_pdf_over_max_pages_with_truncate_processes_only_the_first_n_pages(self):
        """truncate=True (2026-09-10, pedido explícito do usuário): em vez
        de recusar o documento inteiro, processa só as primeiras
        `max_pages` -- confirmado que são de fato as PRIMEIRAS (ordem
        preservada), não um subconjunto qualquer."""
        pdf_bytes = _build_pdf_bytes([f"conteúdo da página {i}" for i in range(5)])

        documents = load_pdf(pdf_bytes, "grande.pdf", max_pages=3, truncate=True)

        assert len(documents) == 3
        assert "página 0" in documents[0].text
        assert "página 1" in documents[1].text
        assert "página 2" in documents[2].text

    def test_pdf_within_max_pages_with_truncate_is_unaffected(self):
        """truncate=True não corta nada quando o documento já está dentro
        do limite -- o flag só importa quando o limite é excedido."""
        pdf_bytes = _build_pdf_bytes(["única página"])

        documents = load_pdf(pdf_bytes, "pequeno.pdf", max_pages=300, truncate=True)

        assert len(documents) == 1


class TestLoadDocx:
    def test_extracts_one_document_per_heading_section_plus_tables(self):
        docx_bytes = _build_docx_bytes()

        documents = load_docx(docx_bytes, "guia.docx", max_pages=300)

        labels = [d.source_label for d in documents]
        assert "guia.docx - Introdução" in labels
        assert "guia.docx - Configuração" in labels
        assert "guia.docx - Tabelas" in labels

        intro = next(d for d in documents if d.source_label == "guia.docx - Introdução")
        assert "parágrafo introdutório" in intro.text

        config_section = next(d for d in documents if d.source_label == "guia.docx - Configuração")
        assert "Primeiro passo" in config_section.text
        assert "Segundo passo" in config_section.text

        tables = next(d for d in documents if d.source_label == "guia.docx - Tabelas")
        assert "Timeout" in tables.text
        assert "30s" in tables.text

    def test_corrupted_docx_raises_a_clear_error(self):
        with pytest.raises(UnsupportedDocumentError):
            load_docx(b"nao e um docx", "quebrado.docx", max_pages=300)

    def test_empty_docx_raises_a_clear_error(self):
        doc = DocxDocument()
        buffer = io.BytesIO()
        doc.save(buffer)

        with pytest.raises(UnsupportedDocumentError, match="texto extraível"):
            load_docx(buffer.getvalue(), "vazio.docx", max_pages=300)

    def test_docx_over_max_pages_raises_page_limit_exceeded_with_real_numbers(self):
        docx_bytes = _build_docx_bytes()  # 3 seções: Introdução, Configuração, Tabelas

        with pytest.raises(PageLimitExceededError) as exc_info:
            load_docx(docx_bytes, "guia.docx", max_pages=2)

        assert exc_info.value.actual_count == 3
        assert exc_info.value.max_pages == 2
        assert exc_info.value.unit == "seções"

    def test_docx_over_max_pages_with_truncate_keeps_only_the_first_n_sections(self):
        docx_bytes = _build_docx_bytes()  # Introdução, Configuração, Tabelas -- nesta ordem

        documents = load_docx(docx_bytes, "guia.docx", max_pages=2, truncate=True)

        assert len(documents) == 2
        labels = [d.source_label for d in documents]
        assert labels == ["guia.docx - Introdução", "guia.docx - Configuração"]


class TestLoadDocumentDispatcher:
    def test_dispatches_pdf_by_extension(self):
        pdf_bytes = _build_pdf_bytes(["conteúdo"])
        documents = load_document(pdf_bytes, "arquivo.PDF", max_pages=300)  # extensão em maiúsculas
        assert len(documents) == 1

    def test_dispatches_docx_by_extension(self):
        docx_bytes = _build_docx_bytes()
        documents = load_document(docx_bytes, "arquivo.docx", max_pages=300)
        assert len(documents) >= 1

    def test_unsupported_extension_raises_with_the_supported_list(self):
        with pytest.raises(UnsupportedDocumentError, match=r"\.docx.*\.pdf|\.pdf.*\.docx"):
            load_document(b"conteudo", "planilha.xlsx", max_pages=300)

    def test_dispatcher_forwards_truncate_to_the_real_loader(self):
        """Confirma que o dispatcher não engole o parâmetro -- sem isso, a
        API poderia aceitar `truncate=true` do frontend e o loader de
        verdade nunca saber disso, voltando a recusar o documento."""
        pdf_bytes = _build_pdf_bytes([f"página {i}" for i in range(5)])

        documents = load_document(pdf_bytes, "grande.pdf", max_pages=3, truncate=True)

        assert len(documents) == 3


class TestSplitPdfIntoParts:
    """Fase A (2026-09-10, "PROMPT DE EVOLUÇÃO"/"OBSERVAÇÕES FINAIS" --
    divisão inteligente consciente de estrutura, substitui truncamento
    como único caminho pra documento grande)."""

    def test_groups_whole_chapters_and_falls_back_to_page_windows_only_for_an_oversized_chapter(self):
        """4 capítulos (5, 5, 30, 5 páginas) com max_pages=10: os 3
        capítulos pequenos (cabem inteiros dentro do teto) nunca devem
        aparecer espalhados por mais de 1 parte cada; só o capítulo de 30
        páginas (sozinho já maior que o teto) deve disparar o fallback de
        janela de página. Confirma também que nada do documento se perde."""
        pdf_bytes = _build_multi_chapter_pdf_bytes([5, 5, 30, 5])

        parts = split_pdf_into_parts(pdf_bytes, "livro.pdf", max_pages=10)

        total_pages = sum(part.page_count for part in parts)
        assert total_pages == 1 + 5 + 5 + 30 + 5  # capa + 4 capítulos, nada perdido

        for part in parts:
            assert part.page_count <= 10

        parts_per_chapter: dict[str, int] = {}
        for part in parts:
            for chapter in part.chapters:
                parts_per_chapter[chapter] = parts_per_chapter.get(chapter, 0) + 1
        assert parts_per_chapter["CHAPTER 1"] == 1
        assert parts_per_chapter["CHAPTER 2"] == 1
        assert parts_per_chapter["CHAPTER 4"] == 1
        assert parts_per_chapter["CHAPTER 3"] > 1  # fallback disparou só aqui

        # última página do documento (capítulo 4, página 5) aparece em
        # alguma parte -- nada foi descartado no processo de divisão.
        last_part_texts = [doc.text for doc in parts[-1].documents]
        assert any("capítulo 4, página 5" in text for text in last_part_texts)

    def test_page_ranges_are_contiguous_and_cover_the_whole_document(self):
        """`page_start`/`page_end` de cada parte precisam formar uma
        sequência contígua e completa (sem buraco, sem sobreposição) --
        é o que a API vai expor como metadado real (Fase B)."""
        pdf_bytes = _build_multi_chapter_pdf_bytes([5, 5, 30, 5])

        parts = split_pdf_into_parts(pdf_bytes, "livro.pdf", max_pages=10)

        assert parts[0].page_start == 1
        for previous, current in zip(parts, parts[1:]):
            assert current.page_start == previous.page_end + 1
        assert parts[-1].page_end == 1 + 5 + 5 + 30 + 5

    def test_falls_back_to_plain_page_windows_when_no_chapter_label_pattern_is_detected(self):
        """Sem nenhum "CHAPTER N"/"Capítulo N" real no documento, a divisão
        não deve fingir consciência de capítulo -- vira janela de página
        fixa honesta, com `chapters` vazio em toda parte."""
        pdf_bytes = _build_pdf_without_chapter_labels_bytes(page_count=25)

        parts = split_pdf_into_parts(pdf_bytes, "sem-estrutura.pdf", max_pages=10)

        assert [part.page_count for part in parts] == [10, 10, 5]
        assert sum(part.page_count for part in parts) == 25
        for part in parts:
            assert part.chapters == []

    def test_document_within_max_pages_produces_a_single_part(self):
        pdf_bytes = _build_multi_chapter_pdf_bytes([2, 2])

        parts = split_pdf_into_parts(pdf_bytes, "pequeno.pdf", max_pages=300)

        assert len(parts) == 1
        assert parts[0].page_count == 1 + 2 + 2  # capa + 2 capítulos


class TestSplitDocxIntoParts:
    """Equivalente de `TestSplitPdfIntoParts` pro DOCX (Fase A) -- unidade
    de divisão é SUBSEÇÃO real (1 `LoadedDocument` por heading com
    parágrafo próprio), não página; o agrupamento acontece por CAPÍTULO
    (Heading 1), a mesma regra de `_group_documents_by_chapter_within_
    limit` compartilhada com o PDF."""

    def test_groups_whole_chapters_and_falls_back_to_windows_only_for_an_oversized_chapter(self):
        """4 capítulos com 3, 3, 20 e 3 subseções: os 3 capítulos pequenos
        nunca devem aparecer espalhados por mais de 1 parte cada; só o
        capítulo de 20 subseções (sozinho já maior que o teto) deve
        disparar o fallback de janela fixa."""
        docx_bytes = _build_docx_with_chapters_bytes([3, 3, 20, 3])

        parts = split_docx_into_parts(docx_bytes, "guia.docx", max_pages=6)

        total_subsections = sum(part.page_count for part in parts)
        assert total_subsections == 3 + 3 + 20 + 3

        for part in parts:
            assert part.page_count <= 6

        parts_per_chapter: dict[str, int] = {}
        for part in parts:
            for chapter in part.chapters:
                parts_per_chapter[chapter] = parts_per_chapter.get(chapter, 0) + 1
        assert parts_per_chapter["Capítulo 1"] == 1
        assert parts_per_chapter["Capítulo 2"] == 1
        assert parts_per_chapter["Capítulo 4"] == 1
        assert parts_per_chapter["Capítulo 3"] > 1

    def test_section_ranges_are_contiguous_and_1_based(self):
        docx_bytes = _build_docx_with_chapters_bytes([3, 3, 20, 3])

        parts = split_docx_into_parts(docx_bytes, "guia.docx", max_pages=6)

        assert parts[0].section_start == 1
        for previous, current in zip(parts, parts[1:]):
            assert current.section_start == previous.section_end + 1
        assert parts[-1].section_end == 3 + 3 + 20 + 3

    def test_document_within_max_pages_produces_a_single_part(self):
        docx_bytes = _build_docx_with_chapters_bytes([2, 2])

        parts = split_docx_into_parts(docx_bytes, "pequeno.docx", max_pages=300)

        assert len(parts) == 1
        assert parts[0].page_count == 4


class TestSplitDocumentDispatcher:
    def test_dispatches_pdf_by_extension(self):
        pdf_bytes = _build_multi_chapter_pdf_bytes([2, 2])
        parts = split_document_into_parts(pdf_bytes, "livro.PDF", max_pages=300)
        assert len(parts) == 1

    def test_dispatches_docx_by_extension(self):
        docx_bytes = _build_docx_with_chapters_bytes([2, 2])
        parts = split_document_into_parts(docx_bytes, "guia.docx", max_pages=300)
        assert len(parts) == 1

    def test_unsupported_extension_raises_with_the_supported_list(self):
        with pytest.raises(UnsupportedDocumentError, match=r"\.docx.*\.pdf|\.pdf.*\.docx"):
            split_document_into_parts(b"conteudo", "planilha.xlsx", max_pages=300)
